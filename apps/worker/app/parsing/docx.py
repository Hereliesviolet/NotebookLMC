"""DOCX parsing via python-docx."""

import io

from docx import Document

from app.parsing.sanitize import sanitize_text
from app.parsing.types import ParsedSection

_HEADING_STYLES = {"Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"}


def parse_docx(data: bytes) -> list[ParsedSection]:
    document = Document(io.BytesIO(data))
    sections: list[ParsedSection] = []
    current_heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        text = sanitize_text("\n".join(buffer)).strip()
        if text:
            sections.append(ParsedSection(text=text, heading=current_heading, chunk_type="text"))
        buffer.clear()

    for paragraph in document.paragraphs:
        text = sanitize_text(paragraph.text).strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name in _HEADING_STYLES:
            flush()
            current_heading = text
            continue
        buffer.append(text)
    flush()

    for table in document.tables:
        rows = [
            "| " + " | ".join(sanitize_text(cell.text).strip() for cell in row.cells) + " |"
            for row in table.rows
        ]
        if rows:
            sections.append(
                ParsedSection(text="\n".join(rows), heading=current_heading, chunk_type="table")
            )

    return sections
