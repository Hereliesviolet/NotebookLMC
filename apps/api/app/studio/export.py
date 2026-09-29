"""Server-side Word (.docx) / PDF / PNG export for Studio artifacts
(architecture doc §19). Renders a persisted artifact's `content_json` (shape
depends on `type`, see packages/prompts/studio_*_tool_schema.json) into a
docx.Document and into an HTML string (rendered to PDF via WeasyPrint),
styled to match the in-app Studio views
(apps/frontend/components/studio/Studio*View.tsx) and design tokens
(apps/frontend/app/globals.css). Both renderers share the same tiny markdown
parser and per-type block builders so the two output formats stay visually
consistent. Mindmaps don't go through that document pipeline - they're
exported as a PNG rasterized (via cairosvg) from an SVG built with the same
radial layout math as the in-app graph, see `build_mindmap_svg`/
`build_mindmap_png` below.
"""

import colorsys
import html
import math
import re
import textwrap
from datetime import datetime
from io import BytesIO

import cairosvg
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from jinja2 import BaseLoader, Environment
from markupsafe import Markup
from weasyprint import HTML as WeasyHTML

# Design tokens, converted 1:1 from apps/frontend/app/globals.css (--primary
# etc, all HSL) and the Tailwind colors StudioBriefingView.tsx uses for its
# per-section accents. WeasyPrint 69 renders css hsl() pixel-identical to
# the equivalent hex (verified manually), but python-docx's RGBColor only
# accepts hex - hex is used everywhere so both renderers stay in sync.
PRIMARY = "#463ACB"  # hsl(245 58% 51%) --primary
FOREGROUND = "#22242B"  # hsl(224 12% 15%) --foreground
MUTED = "#6B7280"  # hsl(220 9% 46%) --muted-foreground
MUTED_BG = "#F3F4F7"  # hsl(220 20% 96%) --muted
BORDER = "#E2E4E9"  # hsl(220 15% 90%) --border
RED = "#DC2626"  # tailwind red-600 (StudioBriefingView risks)
GREEN = "#059669"  # tailwind emerald-600 (recommended_actions)
AMBER = "#D97706"  # tailwind amber-600 (open_questions)

ARTIFACT_TITLES = {
    "summary": "Zusammenfassung",
    "faq": "Häufig gestellte Fragen (FAQ)",
    "timeline": "Timeline",
    "briefing": "Briefing",
    "mindmap": "Mindmap",
    "infographic": "Infografik",
}

_BRIEFING_SECTIONS = [
    ("key_points", "Kernpunkte", PRIMARY, "●"),
    ("risks", "Risiken", RED, "⚠"),
    ("recommended_actions", "Empfohlene Maßnahmen", GREEN, "✓"),
    ("open_questions", "Offene Fragen", AMBER, "?"),
]


def _tint(hex_color: str, target_lightness: float = 0.94) -> str:
    """Light card/section background tied to `hex_color`'s hue, at a fixed
    target HSL lightness rather than a fixed blend-toward-white factor - a
    fixed factor makes colors that start out bright (e.g. the indigo
    primary) end up nearly indistinguishable from white while saturated
    colors (red/green/amber) stay clearly visible; a fixed target lightness
    keeps all four accents visually consistent.
    """
    r, g, b = (int(hex_color[i : i + 2], 16) / 255 for i in (1, 3, 5))
    h, _l, s = colorsys.rgb_to_hls(r, g, b)
    r2, g2, b2 = colorsys.hls_to_rgb(h, target_lightness, s)
    return f"#{round(r2 * 255):02X}{round(g2 * 255):02X}{round(b2 * 255):02X}"


PRIMARY_TINT = _tint(PRIMARY)
RED_TINT = _tint(RED)
GREEN_TINT = _tint(GREEN)
AMBER_TINT = _tint(AMBER)

_ACCENT_TINTS = {PRIMARY: PRIMARY_TINT, RED: RED_TINT, GREEN: GREEN_TINT, AMBER: AMBER_TINT}


# --- tiny markdown parsing shared between the docx and HTML renderers ---
# Just enough for Sonnet's typical summary_markdown output (ATX headings,
# **bold**/*italic*/`code`, bullet/numbered lists, blockquotes, fenced code,
# pipe tables) - not a full CommonMark implementation.

_INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")
_TABLE_SEPARATOR = re.compile(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?$")


def _parse_inline(text: str) -> list[tuple[str, str]]:
    tokens: list[tuple[str, str]] = []
    for part in _INLINE_PATTERN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            tokens.append((part[2:-2], "bold"))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            tokens.append((part[1:-1], "code"))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            tokens.append((part[1:-1], "italic"))
        else:
            tokens.append((part, "plain"))
    return tokens


def _split_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def _parse_markdown_blocks(markdown_text: str) -> list[dict]:
    blocks: list[dict] = []
    lines = markdown_text.splitlines()
    list_buffer: list[str] = []
    list_type: str | None = None
    quote_buffer: list[str] = []

    def flush_list() -> None:
        nonlocal list_buffer, list_type
        if list_buffer:
            blocks.append({"type": list_type, "items": list_buffer})
        list_buffer = []
        list_type = None

    def flush_quote() -> None:
        nonlocal quote_buffer
        if quote_buffer:
            blocks.append({"type": "blockquote", "text": " ".join(quote_buffer)})
        quote_buffer = []

    in_code = False
    code_lines: list[str] = []
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        if stripped.startswith("```"):
            if in_code:
                blocks.append({"type": "code", "text": "\n".join(code_lines)})
                code_lines = []
                in_code = False
            else:
                flush_list()
                flush_quote()
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(lines[i])
            i += 1
            continue

        quote_match = re.match(r"^>\s?(.*)$", stripped)
        if quote_match:
            flush_list()
            quote_buffer.append(quote_match.group(1))
            i += 1
            continue
        flush_quote()

        heading_match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading_match:
            flush_list()
            blocks.append(
                {
                    "type": "heading",
                    "level": len(heading_match.group(1)),
                    "text": heading_match.group(2).strip(),
                }
            )
            i += 1
            continue

        if stripped in ("", "---", "***"):
            flush_list()
            i += 1
            continue

        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet_match:
            if list_type != "bullet_list":
                flush_list()
                list_type = "bullet_list"
            list_buffer.append(bullet_match.group(1))
            i += 1
            continue

        numbered_match = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if numbered_match:
            if list_type != "numbered_list":
                flush_list()
                list_type = "numbered_list"
            list_buffer.append(numbered_match.group(1))
            i += 1
            continue

        if (
            stripped.startswith("|")
            and i + 1 < len(lines)
            and _TABLE_SEPARATOR.match(lines[i + 1].strip())
        ):
            flush_list()
            header = _split_table_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_table_row(lines[i]))
                i += 1
            blocks.append({"type": "table", "header": header, "rows": rows})
            continue

        if stripped.startswith("|"):
            flush_list()
            table_lines = [stripped]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            blocks.append({"type": "code", "text": "\n".join(table_lines)})
            continue

        flush_list()
        blocks.append({"type": "paragraph", "text": stripped})
        i += 1

    flush_list()
    flush_quote()
    if in_code and code_lines:
        blocks.append({"type": "code", "text": "\n".join(code_lines)})
    return blocks


# --- docx low-level helpers (borders/shading aren't exposed by python-docx) ---


def _set_paragraph_box(
    paragraph,
    fill_hex: str | None = None,
    border_hex: str | None = None,
    top: bool = True,
    bottom: bool = True,
    left: bool = True,
    right: bool = True,
    size: int = 4,
    space: int = 6,
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if fill_hex:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill_hex.lstrip("#"))
        p_pr.append(shd)
    if border_hex:
        p_bdr = OxmlElement("w:pBdr")
        for edge, enabled in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
            if not enabled:
                continue
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:space"), str(space))
            el.set(qn("w:color"), border_hex.lstrip("#"))
            p_bdr.append(el)
        p_pr.append(p_bdr)


def _set_run_shading(run, fill_hex: str) -> None:
    r_pr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex.lstrip("#"))
    r_pr.append(shd)


def _set_cell_shading(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex.lstrip("#"))
    tc_pr.append(shd)


def _add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _add_inline_runs(paragraph, text: str) -> None:
    for token_text, style in _parse_inline(text):
        run = paragraph.add_run(token_text)
        if style == "bold":
            run.bold = True
        elif style == "italic":
            run.italic = True
        elif style == "code":
            run.font.name = "Courier New"


def _add_meta_badge(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    run = paragraph.add_run(f" {text} ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(PRIMARY.lstrip("#"))
    run.bold = True
    _set_run_shading(run, PRIMARY_TINT)


def _add_title(document, artifact_type: str, notebook_title: str, generated_at: datetime) -> None:
    heading = document.add_heading(f"{ARTIFACT_TITLES[artifact_type]}: {notebook_title}", level=0)
    heading.paragraph_format.space_after = Pt(4)
    _set_paragraph_box(
        heading,
        border_hex=PRIMARY,
        top=False,
        left=False,
        right=False,
        bottom=True,
        size=18,
        space=8,
    )
    _add_meta_badge(document, f"Generiert am {generated_at.strftime('%d.%m.%Y %H:%M')} Uhr")


def _setup_page(document, notebook_title: str, include_header_footer: bool = True) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    usable_width = section.page_width - section.left_margin - section.right_margin

    if not include_header_footer:
        return

    header_paragraph = section.header.paragraphs[0]
    run = header_paragraph.add_run(f"NotebookLM Clone · {notebook_title}")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
    footer_run = footer_paragraph.add_run("\t")
    footer_run.font.size = Pt(8)
    prefix_run = footer_paragraph.add_run("Seite ")
    prefix_run.font.size = Pt(8)
    prefix_run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
    _add_field(footer_paragraph, "PAGE")
    mid_run = footer_paragraph.add_run(" / ")
    mid_run.font.size = Pt(8)
    mid_run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
    _add_field(footer_paragraph, "NUMPAGES")


# --- docx: per-block-type rendering ---


def _add_heading_with_accent(document, text: str, level: int) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor.from_string(PRIMARY.lstrip("#"))
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(4)
    _set_paragraph_box(
        heading,
        border_hex=PRIMARY,
        top=False,
        left=False,
        right=False,
        bottom=True,
        size=6,
        space=4,
    )


def _add_blockquote_to_docx(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(16)
    _set_paragraph_box(
        paragraph,
        border_hex=PRIMARY,
        top=False,
        bottom=False,
        right=False,
        left=True,
        size=18,
        space=10,
    )
    _add_inline_runs(paragraph, text)
    for run in paragraph.runs:
        run.italic = True
        run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))


def _add_table_to_docx(document, block: dict) -> None:
    header, rows = block["header"], block["rows"]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        _set_cell_shading(cell, PRIMARY)
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("FFFFFF")
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        if row_index % 2 == 1:
            for cell in cells:
                _set_cell_shading(cell, MUTED_BG)
        for cell, text in zip(cells, row_values):
            _add_inline_runs(cell.paragraphs[0], text)
    document.add_paragraph()


def _markdown_blocks_to_docx(document, blocks: list[dict]) -> None:
    for block in blocks:
        if block["type"] == "heading":
            _add_heading_with_accent(document, block["text"], block["level"])
        elif block["type"] in ("bullet_list", "numbered_list"):
            style = "List Bullet" if block["type"] == "bullet_list" else "List Number"
            for item in block["items"]:
                _add_inline_runs(document.add_paragraph(style=style), item)
        elif block["type"] == "blockquote":
            _add_blockquote_to_docx(document, block["text"])
        elif block["type"] == "table":
            _add_table_to_docx(document, block)
        elif block["type"] == "code":
            run = document.add_paragraph().add_run(block["text"])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        else:
            _add_inline_runs(document.add_paragraph(), block["text"])


def _faq_to_docx(document, content: dict) -> None:
    items = content.get("items") or []
    for item in items:
        q_paragraph = document.add_paragraph()
        _set_paragraph_box(
            q_paragraph, fill_hex=PRIMARY_TINT, border_hex=BORDER, bottom=False, size=4, space=8
        )
        q_paragraph.paragraph_format.space_before = Pt(4)
        q_paragraph.paragraph_format.space_after = Pt(2)
        q_run = q_paragraph.add_run("Q ")
        q_run.bold = True
        q_run.font.color.rgb = RGBColor.from_string(PRIMARY.lstrip("#"))
        _set_run_shading(q_run, PRIMARY_TINT)
        text_run = q_paragraph.add_run(item.get("question", ""))
        text_run.bold = True
        text_run.font.color.rgb = RGBColor.from_string(PRIMARY.lstrip("#"))

        a_paragraph = document.add_paragraph()
        a_paragraph.paragraph_format.left_indent = Pt(16)
        source_ids = item.get("source_ids") or []
        _set_paragraph_box(
            a_paragraph,
            fill_hex=PRIMARY_TINT,
            border_hex=BORDER,
            top=False,
            bottom=(not source_ids),
            size=4,
            space=8,
        )
        _add_inline_runs(a_paragraph, item.get("answer", ""))

        if source_ids:
            meta_paragraph = document.add_paragraph()
            meta_paragraph.paragraph_format.left_indent = Pt(16)
            _set_paragraph_box(
                meta_paragraph, fill_hex=PRIMARY_TINT, border_hex=BORDER, top=False, size=4, space=8
            )
            meta_run = meta_paragraph.add_run("Quellen: " + ", ".join(source_ids))
            meta_run.italic = True
            meta_run.font.size = Pt(8)
            meta_run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))

        document.add_paragraph().paragraph_format.space_after = Pt(2)


def _sorted_timeline_events(content: dict) -> list[dict]:
    events = content.get("events") or []
    dated = [e for e in events if e.get("date")]
    undated = [e for e in events if not e.get("date")]
    dated.sort(key=lambda e: e["date"])
    return dated + undated


def _timeline_to_docx(document, content: dict) -> None:
    events = _sorted_timeline_events(content)
    if not events:
        document.add_paragraph("Keine Ereignisse gefunden.")
        return

    table = document.add_table(rows=0, cols=2)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)
    table.columns[0].width = Cm(0.9)
    table.columns[1].width = Cm(14.5)

    for event in events:
        row = table.add_row().cells
        dot_run = row[0].paragraphs[0].add_run("●")
        dot_run.font.color.rgb = RGBColor.from_string(PRIMARY.lstrip("#"))
        dot_run.font.size = Pt(13)

        label_paragraph = row[1].paragraphs[0]
        label_run = label_paragraph.add_run(
            event.get("date_label") or event.get("date") or "Unklar"
        )
        label_run.bold = True
        label_run.font.color.rgb = RGBColor.from_string(PRIMARY.lstrip("#"))

        desc_paragraph = row[1].add_paragraph()
        desc_paragraph.paragraph_format.space_before = Pt(2)
        _add_inline_runs(desc_paragraph, event.get("description", ""))

        quote = event.get("quote")
        if quote:
            quote_paragraph = row[1].add_paragraph()
            quote_paragraph.paragraph_format.left_indent = Pt(10)
            _set_paragraph_box(
                quote_paragraph,
                border_hex=BORDER,
                top=False,
                bottom=False,
                right=False,
                size=10,
                space=6,
            )
            quote_run = quote_paragraph.add_run(f"„{quote}“")
            quote_run.italic = True
            quote_run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
            quote_run.font.size = Pt(9.5)

        source_paragraph = row[1].add_paragraph()
        source_paragraph.paragraph_format.space_after = Pt(10)
        source_run = source_paragraph.add_run(f"Quelle: {event.get('source_id', '')}")
        source_run.italic = True
        source_run.font.size = Pt(8)
        source_run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))


def _briefing_to_docx(document, content: dict) -> None:
    intro = document.add_paragraph()
    _add_inline_runs(intro, content.get("summary", ""))
    intro.paragraph_format.space_after = Pt(14)

    for field, title, color, icon in _BRIEFING_SECTIONS:
        items = content.get(field) or []
        tint = _ACCENT_TINTS[color]

        box_kwargs = dict(border_hex=color, top=False, bottom=False, right=False, size=24, space=10)

        heading_paragraph = document.add_paragraph()
        _set_paragraph_box(heading_paragraph, fill_hex=tint, **box_kwargs)
        heading_paragraph.paragraph_format.space_before = Pt(6)
        heading_paragraph.paragraph_format.space_after = Pt(2)
        heading_run = heading_paragraph.add_run(f"{icon}  {title}")
        heading_run.bold = True
        heading_run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))

        if not items:
            empty_paragraph = document.add_paragraph()
            _set_paragraph_box(empty_paragraph, fill_hex=tint, **box_kwargs)
            empty_run = empty_paragraph.add_run("– keine –")
            empty_run.italic = True
            empty_run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
        else:
            for item in items:
                item_paragraph = document.add_paragraph()
                _set_paragraph_box(item_paragraph, fill_hex=tint, **box_kwargs)
                bullet_run = item_paragraph.add_run(f"{icon}  ")
                bullet_run.bold = True
                bullet_run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
                _add_inline_runs(item_paragraph, item)

        document.add_paragraph().paragraph_format.space_after = Pt(2)


def _infographic_to_docx(document, content: dict) -> None:
    """Simple text fallback for the poster layout (a CSS-grid poster isn't
    reproducible in Word) - headline/subheadline as a plain title, stats as
    a compact line, sections as a numbered list with bold titles.
    """
    headline = document.add_heading(content.get("headline", ""), level=0)
    for run in headline.runs:
        run.font.color.rgb = RGBColor.from_string(PRIMARY.lstrip("#"))
    headline.paragraph_format.space_after = Pt(2)

    subheadline = content.get("subheadline")
    if subheadline:
        sub_paragraph = document.add_paragraph()
        sub_run = sub_paragraph.add_run(subheadline)
        sub_run.italic = True
        sub_run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
    document.add_paragraph().paragraph_format.space_after = Pt(4)

    stats = content.get("stats") or []
    if stats:
        stats_paragraph = document.add_paragraph()
        _set_paragraph_box(
            stats_paragraph, fill_hex=PRIMARY_TINT, border_hex=BORDER, size=4, space=8
        )
        for index, stat in enumerate(stats):
            if index > 0:
                sep_run = stats_paragraph.add_run("    ·    ")
                sep_run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
            value_run = stats_paragraph.add_run(f"{stat.get('value', '')} ")
            value_run.bold = True
            value_run.font.color.rgb = RGBColor.from_string(PRIMARY.lstrip("#"))
            label_run = stats_paragraph.add_run(stat.get("label", ""))
            label_run.font.size = Pt(9)
            label_run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
        document.add_paragraph().paragraph_format.space_after = Pt(4)

    for index, section in enumerate(content.get("sections") or [], start=1):
        section_paragraph = document.add_paragraph()
        _set_paragraph_box(
            section_paragraph, fill_hex=MUTED_BG, border_hex=BORDER, bottom=False, size=4, space=8
        )
        section_paragraph.paragraph_format.space_before = Pt(4)
        badge_run = section_paragraph.add_run(f" {index} ")
        badge_run.bold = True
        badge_run.font.color.rgb = RGBColor.from_string("FFFFFF")
        _set_run_shading(badge_run, PRIMARY)
        title_run = section_paragraph.add_run(f"  {section.get('title', '')}")
        title_run.bold = True
        title_run.font.color.rgb = RGBColor.from_string(PRIMARY.lstrip("#"))

        body_paragraph = document.add_paragraph()
        _set_paragraph_box(
            body_paragraph, fill_hex=MUTED_BG, border_hex=BORDER, top=False, size=4, space=8
        )
        _add_inline_runs(body_paragraph, section.get("body", ""))
        document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(
    artifact_type: str, notebook_title: str, content: dict, generated_at: datetime
) -> Document:
    document = Document()
    is_infographic = artifact_type == "infographic"
    _setup_page(document, notebook_title, include_header_footer=not is_infographic)
    if not is_infographic:
        _add_title(document, artifact_type, notebook_title, generated_at)

    if artifact_type == "summary":
        _markdown_blocks_to_docx(
            document, _parse_markdown_blocks(content.get("summary_markdown", ""))
        )
    elif artifact_type == "faq":
        _faq_to_docx(document, content)
    elif artifact_type == "timeline":
        _timeline_to_docx(document, content)
    elif artifact_type == "briefing":
        _briefing_to_docx(document, content)
    elif artifact_type == "infographic":
        _infographic_to_docx(document, content)
    else:
        raise ValueError(f"Unknown studio artifact type: {artifact_type!r}")
    return document


def render_docx_bytes(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- HTML / PDF rendering ---


def _inline_to_html(text: str) -> str:
    parts = []
    for token_text, style in _parse_inline(text):
        escaped = html.escape(token_text)
        if style == "bold":
            parts.append(f"<strong>{escaped}</strong>")
        elif style == "italic":
            parts.append(f"<em>{escaped}</em>")
        elif style == "code":
            parts.append(f"<code>{escaped}</code>")
        else:
            parts.append(escaped)
    return "".join(parts)


def _table_block_to_html(block: dict) -> str:
    header_html = "".join(f"<th>{_inline_to_html(c)}</th>" for c in block["header"])
    rows_html = "".join(
        "<tr>" + "".join(f"<td>{_inline_to_html(c)}</td>" for c in row) + "</tr>"
        for row in block["rows"]
    )
    return f"<table class='doc-table'><thead><tr>{header_html}</tr></thead><tbody>{rows_html}</tbody></table>"


def _markdown_blocks_to_html(blocks: list[dict]) -> str:
    parts = []
    for block in blocks:
        if block["type"] == "heading":
            level = block["level"] + 1  # +1: document <h1> is the artifact title
            parts.append(f"<h{level}>{_inline_to_html(block['text'])}</h{level}>")
        elif block["type"] in ("bullet_list", "numbered_list"):
            tag = "ul" if block["type"] == "bullet_list" else "ol"
            items = "".join(f"<li>{_inline_to_html(i)}</li>" for i in block["items"])
            parts.append(f"<{tag}>{items}</{tag}>")
        elif block["type"] == "blockquote":
            parts.append(f"<blockquote>{_inline_to_html(block['text'])}</blockquote>")
        elif block["type"] == "table":
            parts.append(_table_block_to_html(block))
        elif block["type"] == "code":
            parts.append(f"<pre>{html.escape(block['text'])}</pre>")
        else:
            parts.append(f"<p>{_inline_to_html(block['text'])}</p>")
    return "\n".join(parts)


def _faq_to_html(content: dict) -> str:
    items = content.get("items") or []
    if not items:
        return "<p class='meta'>Keine Fragen gefunden.</p>"
    parts = []
    for item in items:
        source_ids = item.get("source_ids") or []
        meta = (
            f"<p class='faq-meta'>Quellen: {html.escape(', '.join(source_ids))}</p>"
            if source_ids
            else ""
        )
        parts.append(
            "<div class='faq-card'>"
            "<div class='faq-q'><span class='faq-q-badge'>Q</span>"
            f"<span>{_inline_to_html(item.get('question', ''))}</span></div>"
            f"<p class='faq-a'>{_inline_to_html(item.get('answer', ''))}</p>"
            f"{meta}"
            "</div>"
        )
    return "\n".join(parts)


def _timeline_to_html(content: dict) -> str:
    events = _sorted_timeline_events(content)
    if not events:
        return "<p class='meta'>Keine Ereignisse gefunden.</p>"
    parts = ["<div class='timeline'>"]
    for event in events:
        label = event.get("date_label") or event.get("date") or "Unklar"
        quote = event.get("quote")
        quote_html = f"<p class='timeline-quote'>„{_inline_to_html(quote)}“</p>" if quote else ""
        parts.append(
            "<div class='timeline-item'><span class='timeline-dot'></span>"
            f"<p class='timeline-label'>{html.escape(str(label))}</p>"
            f"<p class='timeline-desc'>{_inline_to_html(event.get('description', ''))}</p>"
            f"{quote_html}"
            f"<p class='timeline-source'>Quelle: {html.escape(str(event.get('source_id', '')))}</p>"
            "</div>"
        )
    parts.append("</div>")
    return "\n".join(parts)


def _briefing_to_html(content: dict) -> str:
    parts = [f"<p>{_inline_to_html(content.get('summary', ''))}</p>"]
    for field, title, color, icon in _BRIEFING_SECTIONS:
        items = content.get(field) or []
        tint = _ACCENT_TINTS[color]
        body = (
            "<p class='briefing-empty'>– keine –</p>"
            if not items
            else "<ul>" + "".join(f"<li>{_inline_to_html(i)}</li>" for i in items) + "</ul>"
        )
        parts.append(
            f"<div class='briefing-section' style='background:{tint}; border-left-color:{color};'>"
            f"<h2 style='color:{color};'><span class='briefing-icon' style='background:{color};'>{icon}</span>{title}</h2>"
            f"{body}"
            "</div>"
        )
    return "\n".join(parts)


def _build_infographic_html(content: dict) -> str:
    """Poster layout (not a flowing document like the other types): a
    PRIMARY-colored header band, an optional stat-box row, and a 2-column
    card grid with numbered badges (analog to the `.faq-q-badge` pattern
    above) - mirrors StudioInfographicView.tsx 1:1 so the in-app preview and
    the exported PDF stay visually consistent.
    """
    stats = content.get("stats") or []
    stats_html = ""
    if stats:
        stat_items = "".join(
            "<div class='infographic-stat'>"
            f"<div class='infographic-stat-value'>{html.escape(stat.get('value', ''))}</div>"
            f"<div class='infographic-stat-label'>{html.escape(stat.get('label', ''))}</div>"
            "</div>"
            for stat in stats
        )
        stats_html = f"<div class='infographic-stats'>{stat_items}</div>"

    cards_html = "".join(
        "<div class='infographic-card'>"
        "<div class='infographic-card-header'>"
        f"<span class='infographic-badge'>{index}</span>"
        f"<h3 class='infographic-card-title'>{html.escape(section.get('title', ''))}</h3>"
        "</div>"
        f"<p class='infographic-card-body'>{html.escape(section.get('body', ''))}</p>"
        "</div>"
        for index, section in enumerate(content.get("sections") or [], start=1)
    )

    return (
        "<div class='infographic-header'>"
        f"<h1 class='infographic-headline'>{html.escape(content.get('headline', ''))}</h1>"
        f"<p class='infographic-subheadline'>{html.escape(content.get('subheadline', ''))}</p>"
        "</div>"
        f"{stats_html}"
        f"<div class='infographic-grid'>{cards_html}</div>"
    )


def _css_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace('"', '\\"')


_HTML_TEMPLATE = Environment(loader=BaseLoader(), autoescape=True).from_string(
    """<!DOCTYPE html>
<html lang="de">
<head>
<meta charset="utf-8">
<title>{{ title }}</title>
<style>
  @page {
    size: A4;
    margin: 2.6cm 2cm 2.4cm 2cm;
    {%% if show_chrome %%}
    @top-left { content: "NotebookLM Clone · {{ notebook_title_css }}"; font-size: 8pt; color: %(muted)s; }
    @bottom-right { content: "Seite " counter(page) " / " counter(pages); font-size: 8pt; color: %(muted)s; }
    {%% endif %%}
  }
  body { font-family: 'Liberation Sans', Arial, sans-serif; color: %(foreground)s; font-size: 11pt; line-height: 1.6; }
  .doc-title { font-size: 22pt; font-weight: 700; margin: 0 0 10pt 0; padding-bottom: 10pt; border-bottom: 3px solid %(primary)s; }
  .doc-meta { display: inline-block; margin: 12pt 0 22pt 0; padding: 3pt 12pt; border-radius: 10pt; background: %(primary_tint)s; color: %(primary)s; font-size: 9pt; font-weight: 600; }
  h2 { font-size: 14pt; color: %(primary)s; border-bottom: 1.5px solid %(primary)s; padding-bottom: 3pt; margin-top: 22pt; }
  h3, h4 { font-size: 12pt; color: %(primary)s; border-bottom: 1px solid %(primary_tint)s; padding-bottom: 2pt; margin-top: 14pt; }
  p { margin: 7pt 0; }
  ul, ol { margin: 6pt 0; padding-left: 20pt; }
  blockquote { border-left: 3px solid %(primary)s; margin: 10pt 0; padding: 5pt 12pt; color: %(muted)s; font-style: italic; background: %(muted_bg)s; }
  code, pre { font-family: 'Courier New', monospace; font-size: 9.5pt; }
  pre { background: %(muted_bg)s; padding: 8pt; white-space: pre-wrap; border-radius: 4pt; }
  .meta { color: %(muted)s; font-size: 9pt; }

  .doc-table { border-collapse: collapse; width: 100%%; margin: 10pt 0; font-size: 10pt; }
  .doc-table th { background: %(primary)s; color: #fff; text-align: left; padding: 5pt 8pt; }
  .doc-table td { border: 1px solid %(border)s; padding: 5pt 8pt; }
  .doc-table tr:nth-child(even) td { background: %(muted_bg)s; }

  .faq-card { border: 1px solid %(border)s; border-radius: 8pt; background: %(primary_tint)s; padding: 12pt 14pt; margin-bottom: 10pt; page-break-inside: avoid; }
  .faq-q { display: flex; gap: 8pt; align-items: flex-start; font-weight: 700; color: %(primary)s; margin-bottom: 6pt; }
  .faq-q-badge { flex: 0 0 auto; width: 16pt; height: 16pt; border-radius: 50%%; background: %(primary)s; color: #fff; font-size: 9pt; text-align: center; line-height: 16pt; }
  .faq-a { margin: 0 0 0 24pt; color: %(foreground)s; }
  .faq-meta { margin: 4pt 0 0 24pt; color: %(muted)s; font-size: 8.5pt; }

  .timeline { position: relative; margin-left: 8pt; padding-left: 20pt; border-left: 2.5px solid %(primary)s; }
  .timeline-item { position: relative; margin-bottom: 18pt; page-break-inside: avoid; }
  .timeline-dot { position: absolute; left: -25.5pt; top: 3pt; width: 9pt; height: 9pt; border-radius: 50%%; background: %(primary)s; border: 2px solid #fff; box-shadow: 0 0 0 1px %(primary)s; }
  .timeline-label { font-weight: 700; color: %(primary)s; font-size: 10pt; margin: 0 0 2pt 0; }
  .timeline-desc { margin: 2pt 0; }
  .timeline-quote { border-left: 2px solid %(border)s; margin: 4pt 0 0 0; padding-left: 8pt; color: %(muted)s; font-style: italic; font-size: 9.5pt; }
  .timeline-source { margin: 3pt 0 0 0; color: %(muted)s; font-size: 8.5pt; }

  .briefing-section { border-radius: 8pt; border-left: 4px solid; padding: 12pt 14pt; margin-bottom: 12pt; page-break-inside: avoid; }
  .briefing-section h2 { border: none; margin: 0 0 8pt 0; padding: 0; display: flex; align-items: center; gap: 7pt; font-size: 12.5pt; }
  .briefing-icon { display: inline-flex; align-items: center; justify-content: center; width: 16pt; height: 16pt; border-radius: 50%%; color: #fff; font-size: 9pt; }
  .briefing-section ul { margin: 0; padding-left: 18pt; }
  .briefing-section li { margin-bottom: 4pt; }
  .briefing-empty { color: %(muted)s; font-style: italic; margin: 0; }

  .infographic-header { background: %(primary)s; border-radius: 10pt; padding: 22pt 24pt; margin: 0 0 16pt 0; }
  .infographic-headline { margin: 0; font-size: 22pt; font-weight: 700; color: #fff; }
  .infographic-subheadline { margin: 6pt 0 0 0; font-size: 11pt; color: #fff; opacity: 0.85; }
  .infographic-stats { display: flex; gap: 10pt; margin-bottom: 16pt; }
  .infographic-stat { flex: 1; text-align: center; border: 1px solid %(border)s; border-radius: 8pt; background: %(primary_tint)s; padding: 10pt 8pt; }
  .infographic-stat-value { font-size: 16pt; font-weight: 700; color: %(primary)s; }
  .infographic-stat-label { margin-top: 2pt; font-size: 8.5pt; color: %(muted)s; }
  .infographic-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 12pt; }
  .infographic-card { border: 1px solid %(border)s; border-radius: 8pt; padding: 12pt 14pt; background: #fff; page-break-inside: avoid; }
  .infographic-card-header { display: flex; align-items: center; gap: 8pt; margin-bottom: 6pt; }
  .infographic-badge { flex: 0 0 auto; width: 18pt; height: 18pt; border-radius: 50%%; background: %(primary)s; color: #fff; font-size: 10pt; font-weight: 700; text-align: center; line-height: 18pt; }
  .infographic-card-title { margin: 0; font-size: 11.5pt; color: %(primary)s; font-weight: 700; }
  .infographic-card-body { margin: 0; font-size: 10pt; color: %(foreground)s; }
</style>
</head>
<body>
{%% if show_chrome %%}
<h1 class="doc-title">{{ title }}</h1>
<p class="doc-meta">{{ subtitle }}</p>
{%% endif %%}
{{ body | safe }}
</body>
</html>"""
    % {
        "primary": PRIMARY,
        "primary_tint": PRIMARY_TINT,
        "muted": MUTED,
        "muted_bg": MUTED_BG,
        "border": BORDER,
        "foreground": FOREGROUND,
    }
)


def build_html(
    artifact_type: str, notebook_title: str, content: dict, generated_at: datetime
) -> str:
    if artifact_type == "summary":
        body = _markdown_blocks_to_html(_parse_markdown_blocks(content.get("summary_markdown", "")))
    elif artifact_type == "faq":
        body = _faq_to_html(content)
    elif artifact_type == "timeline":
        body = _timeline_to_html(content)
    elif artifact_type == "briefing":
        body = _briefing_to_html(content)
    elif artifact_type == "infographic":
        body = _build_infographic_html(content)
    else:
        raise ValueError(f"Unknown studio artifact type: {artifact_type!r}")

    return _HTML_TEMPLATE.render(
        title=f"{ARTIFACT_TITLES[artifact_type]}: {notebook_title}",
        subtitle=f"Generiert am {generated_at.strftime('%d.%m.%Y %H:%M')} Uhr",
        notebook_title_css=Markup(_css_escape(notebook_title)),
        body=body,
        show_chrome=artifact_type != "infographic",
    )


def render_pdf_bytes(html_content: str) -> bytes:
    return WeasyHTML(string=html_content).write_pdf()


# --- Mindmap: SVG / PNG rendering ---
#
# Unlike the other artifact types, the mindmap is exported as a raster image
# of its in-app node graph rather than reflowed into a document. The layout
# math below (radial branch placement, per-branch leaf fan-out) is a 1:1
# port of `buildMindmapGraph()` in
# apps/frontend/components/studio/StudioMindmapView.tsx, so the exported PNG
# matches the on-screen graph. The SVG is the render target because it's
# trivial to build from plain coordinates; cairosvg then rasterizes it.

_MINDMAP_BRANCH_BASE_RADIUS = 260
_MINDMAP_LEAF_BASE_RADIUS = 170

# fill/stroke/text colors mirror the Tailwind classes on Root/Branch/LeafMindmapNode
# in StudioMindmapView.tsx (bg-primary, border-border/bg-card, bg-muted/text-muted-foreground).
_MINDMAP_NODE_STYLES = {
    "root": {
        "fill": PRIMARY,
        "stroke": PRIMARY,
        "stroke_width": 0,
        "text_color": "#FFFFFF",
        "font_weight": 700,
        "font_size": 14,
        "max_w": 220,
        "min_w": 70,
        "pad_x": 16,
        "pad_y": 10,
    },
    "branch": {
        "fill": "#FFFFFF",
        "stroke": BORDER,
        "stroke_width": 1.5,
        "text_color": FOREGROUND,
        "font_weight": 600,
        "font_size": 13,
        "max_w": 190,
        "min_w": 60,
        "pad_x": 14,
        "pad_y": 10,
    },
    "leaf": {
        "fill": MUTED_BG,
        "stroke": MUTED_BG,
        "stroke_width": 0,
        "text_color": MUTED,
        "font_weight": 400,
        "font_size": 11.5,
        "max_w": 170,
        "min_w": 50,
        "pad_x": 12,
        "pad_y": 8,
    },
}


def _mindmap_box_intersection(
    cx: float, cy: float, half_w: float, half_h: float, tx: float, ty: float
) -> tuple[float, float]:
    """Point where the segment from (cx, cy) to (tx, ty) crosses the boundary
    of the axis-aligned box centered on (cx, cy) - used so edges touch node
    boxes at whichever side actually faces the other node (root/branch boxes
    are arranged radially, not stacked top-down), instead of being drawn
    straight into the box and under the label.
    """
    dx, dy = tx - cx, ty - cy
    if dx == 0 and dy == 0:
        return cx, cy
    scales = []
    if dx != 0:
        scales.append(half_w / abs(dx))
    if dy != 0:
        scales.append(half_h / abs(dy))
    scale = min(scales)
    return cx + dx * scale, cy + dy * scale


def build_mindmap_svg(content: dict) -> str:
    root_data = content.get("root") or {}
    branches = root_data.get("children") or []
    branch_count = max(len(branches), 1)
    branch_radius = max(_MINDMAP_BRANCH_BASE_RADIUS, branch_count * 55)

    nodes: dict[str, dict] = {}
    edges: list[tuple[str, str]] = []

    def add_node(node_id: str, kind: str, x: float, y: float, label: str) -> None:
        style = _MINDMAP_NODE_STYLES[kind]
        font_size = style["font_size"]
        pad_x, pad_y = style["pad_x"], style["pad_y"]
        avg_char_width = font_size * 0.62
        line_height = font_size * 1.3
        max_chars = max(6, int((style["max_w"] - 2 * pad_x) // avg_char_width))
        lines = textwrap.wrap(label or "", width=max_chars) or [""]
        longest_line = max(len(line) for line in lines)
        width = min(style["max_w"], max(style["min_w"], longest_line * avg_char_width + 2 * pad_x))
        height = 2 * pad_y + line_height * len(lines)
        nodes[node_id] = {
            "kind": kind,
            "x": x,
            "y": y,
            "lines": lines,
            "width": width,
            "height": height,
            "font_size": font_size,
            "line_height": line_height,
        }

    add_node("root", "root", 0.0, 0.0, root_data.get("label", ""))

    for branch_index, branch in enumerate(branches):
        branch_angle = (branch_index / branch_count) * 2 * math.pi - math.pi / 2
        branch_x = math.cos(branch_angle) * branch_radius
        branch_y = math.sin(branch_angle) * branch_radius
        branch_id = f"branch-{branch_index}"
        add_node(branch_id, "branch", branch_x, branch_y, branch.get("label", ""))
        edges.append(("root", branch_id))

        leaves = branch.get("children") or []
        leaf_count = len(leaves)
        if leaf_count == 0:
            continue

        fan_spread = min(math.pi * 0.85, 0.5 + leaf_count * 0.45)
        leaf_radius = max(_MINDMAP_LEAF_BASE_RADIUS, leaf_count * 26)

        for leaf_index, leaf in enumerate(leaves):
            if leaf_count == 1:
                leaf_angle = branch_angle
            else:
                leaf_angle = (
                    branch_angle - fan_spread / 2 + (fan_spread * leaf_index) / (leaf_count - 1)
                )
            leaf_x = branch_x + math.cos(leaf_angle) * leaf_radius
            leaf_y = branch_y + math.sin(leaf_angle) * leaf_radius
            leaf_id = f"{branch_id}-leaf-{leaf_index}"
            add_node(leaf_id, "leaf", leaf_x, leaf_y, leaf.get("label", ""))
            edges.append((branch_id, leaf_id))

    padding = 40
    min_x = min(n["x"] - n["width"] / 2 for n in nodes.values())
    max_x = max(n["x"] + n["width"] / 2 for n in nodes.values())
    min_y = min(n["y"] - n["height"] / 2 for n in nodes.values())
    max_y = max(n["y"] + n["height"] / 2 for n in nodes.values())
    offset_x, offset_y = padding - min_x, padding - min_y
    canvas_width = (max_x - min_x) + 2 * padding
    canvas_height = (max_y - min_y) + 2 * padding

    def shifted(node: dict) -> tuple[float, float]:
        return node["x"] + offset_x, node["y"] + offset_y

    svg_parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {canvas_width:.1f} {canvas_height:.1f}" '
        f'width="{canvas_width:.1f}" height="{canvas_height:.1f}" font-family="Liberation Sans, Arial, sans-serif">',
        f'<rect x="0" y="0" width="{canvas_width:.1f}" height="{canvas_height:.1f}" fill="#FFFFFF" />',
    ]

    edge_parts = []
    for source_id, target_id in edges:
        source, target = nodes[source_id], nodes[target_id]
        sx, sy = shifted(source)
        tx, ty = shifted(target)
        sx2, sy2 = _mindmap_box_intersection(
            sx, sy, source["width"] / 2, source["height"] / 2, tx, ty
        )
        tx2, ty2 = _mindmap_box_intersection(
            tx, ty, target["width"] / 2, target["height"] / 2, sx, sy
        )
        edge_parts.append(
            f'<line x1="{sx2:.1f}" y1="{sy2:.1f}" x2="{tx2:.1f}" y2="{ty2:.1f}" '
            f'stroke="{BORDER}" stroke-width="1.5" />'
        )
    svg_parts.append(f"<g>{''.join(edge_parts)}</g>")

    # Rects and text are each drawn in their own pass (all boxes, then all
    # labels on top) rather than per-node - closely packed leaf fans can have
    # slightly overlapping boxes (the label-width estimate below is
    # approximate), and interleaving would let a later box silently paint
    # over an earlier label's edge.
    rect_parts = []
    text_parts = []
    for node in nodes.values():
        style = _MINDMAP_NODE_STYLES[node["kind"]]
        cx, cy = shifted(node)
        rect_x, rect_y = cx - node["width"] / 2, cy - node["height"] / 2
        rect_parts.append(
            f'<rect x="{rect_x:.1f}" y="{rect_y:.1f}" width="{node["width"]:.1f}" height="{node["height"]:.1f}" '
            f'rx="8" fill="{style["fill"]}" stroke="{style["stroke"]}" stroke-width="{style["stroke_width"]}" />'
        )
        text_height = node["line_height"] * len(node["lines"])
        first_baseline = cy - text_height / 2 + node["font_size"] * 0.85
        for line_index, line in enumerate(node["lines"]):
            baseline_y = first_baseline + line_index * node["line_height"]
            text_parts.append(
                f'<text x="{cx:.1f}" y="{baseline_y:.1f}" text-anchor="middle" '
                f'font-size="{node["font_size"]}" font-weight="{style["font_weight"]}" '
                f'fill="{style["text_color"]}">{html.escape(line)}</text>'
            )
    svg_parts.append(f"<g>{''.join(rect_parts)}</g>")
    svg_parts.append(f"<g>{''.join(text_parts)}</g>")
    svg_parts.append("</svg>")
    return "".join(svg_parts)


def build_mindmap_png(content: dict) -> bytes:
    svg = build_mindmap_svg(content)
    return cairosvg.svg2png(bytestring=svg.encode("utf-8"), background_color="white", scale=2.0)


def safe_filename(base: str, extension: str) -> str:
    ascii_base = base.encode("ascii", "ignore").decode("ascii")
    ascii_base = re.sub(r"[^A-Za-z0-9._ -]", "_", ascii_base).strip(" ._") or "export"
    return f"{ascii_base}.{extension}"
